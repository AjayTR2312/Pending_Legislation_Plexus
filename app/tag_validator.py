"""
Tag Validator — Automated comparison of Claude output vs PyMuPDF ground truth
==============================================================================

Compares Claude's ``[A>…<A]`` / ``[D>…<D]`` tagged text against the PyMuPDF
JSON report (which detects formatting directly from the PDF structure).

Produces an **Excel report** (.xlsx) with:
  • Sheet 1 — **Span Validation**: one row per PyMuPDF-detected span
  • Sheet 2 — **False Positives**: Claude tags with no corresponding PyMuPDF span
  • Sheet 3 — **Summary**: precision, recall, grade, counts

Usage (standalone)::

    python tag_validator.py  <claude_output.txt>  <pymupdf_report.json>  [-o report.xlsx]

Usage (programmatic)::

    from tag_validator import validate_conversion
    validate_conversion(claude_text, pymupdf_report, "validation.xlsx")
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------
# Minimum similarity ratio for a fuzzy match
FUZZY_THRESHOLD = 0.75

# Format type → expected tag direction
_FORMAT_TO_TAG = {
    "underline":     "add",
    "italic":        "add",
    "bold":          "add",
    "strikethrough": "delete",
}


# ---------------------------------------------------------------------------
# Data structures
# ---------------------------------------------------------------------------
@dataclass
class ClaudeTag:
    """A single tagged span extracted from Claude's output."""
    text: str               # raw text inside the tag
    tag_type: str           # "add" or "delete"
    start_pos: int = 0      # character position in Claude output
    end_pos: int = 0
    matched: bool = False   # set True when matched to a PyMuPDF span


@dataclass
class ValidationRow:
    """One row in the Span Validation sheet."""
    page: int
    bill_line: Optional[int]
    text: str
    format_types: str           # comma-joined: "underline", "strikethrough,italic"
    expected_tag: str           # "[A>...<A]" or "[D>...<D]"
    context: str
    tag_found: str              # "✅ Yes", "❌ No", "⚠️ Partial"
    text_match: str             # "✅ Full", "⚠️ Partial", "❌ Missing"
    missing_chars: str          # characters in PyMuPDF but not in Claude tag
    extra_chars: str            # characters in Claude tag but not in PyMuPDF
    severity: str               # "OK", "Minor", "Critical"
    notes: str


@dataclass
class FalsePositiveRow:
    """One row in the False Positives sheet."""
    tag_type: str               # "[A>...<A]" or "[D>...<D]"
    tagged_text: str
    position_in_output: int     # char offset in Claude output
    notes: str


# ---------------------------------------------------------------------------
# Parse Claude tags
# ---------------------------------------------------------------------------
_RE_ADD_TAG = re.compile(r"\[A>(.*?)<A\]", re.DOTALL)
_RE_DEL_TAG = re.compile(r"\[D>(.*?)<D\]", re.DOTALL)


def parse_claude_tags(text: str) -> List[ClaudeTag]:
    """
    Extract all ``[A>…<A]`` and ``[D>…<D]`` spans from Claude's output.

    Returns a list of :class:`ClaudeTag` sorted by position.
    """
    tags: List[ClaudeTag] = []
    for m in _RE_ADD_TAG.finditer(text):
        tags.append(ClaudeTag(
            text=m.group(1), tag_type="add",
            start_pos=m.start(), end_pos=m.end(),
        ))
    for m in _RE_DEL_TAG.finditer(text):
        tags.append(ClaudeTag(
            text=m.group(1), tag_type="delete",
            start_pos=m.start(), end_pos=m.end(),
        ))
    tags.sort(key=lambda t: t.start_pos)
    return tags


# ---------------------------------------------------------------------------
# Normalisation helpers
# ---------------------------------------------------------------------------

def _normalise(text: str) -> str:
    """Lowercase, collapse whitespace, strip."""
    return re.sub(r"\s+", " ", text.strip().lower())


def _char_diff(expected: str, actual: str) -> Tuple[str, str]:
    """
    Return (missing_chars, extra_chars) between *expected* and *actual*.

    missing = characters in expected but not in actual
    extra   = characters in actual but not in expected
    """
    sm = SequenceMatcher(None, expected, actual)
    missing_parts = []
    extra_parts = []
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "delete":
            missing_parts.append(expected[i1:i2])
        elif op == "insert":
            extra_parts.append(actual[j1:j2])
        elif op == "replace":
            missing_parts.append(expected[i1:i2])
            extra_parts.append(actual[j1:j2])
    return "".join(missing_parts), "".join(extra_parts)


# ---------------------------------------------------------------------------
# Span matching
# ---------------------------------------------------------------------------

def _find_best_match(
    pymupdf_text: str,
    expected_tag_type: str,
    claude_tags: List[ClaudeTag],
) -> Optional[ClaudeTag]:
    """
    Find the best matching Claude tag for a PyMuPDF span.

    Matching strategy (in priority order):
      1. Exact normalised text match (same tag type)
      2. Claude tag contains PyMuPDF text (Claude tagged more)
      3. PyMuPDF text contains Claude tag (Claude tagged less — partial)
      4. Fuzzy match above threshold

    Returns the best matching ClaudeTag, or None.
    """
    norm_expected = _normalise(pymupdf_text)
    if not norm_expected:
        return None

    candidates = [t for t in claude_tags if t.tag_type == expected_tag_type and not t.matched]

    best_tag: Optional[ClaudeTag] = None
    best_score: float = 0.0

    for ct in candidates:
        norm_actual = _normalise(ct.text)
        if not norm_actual:
            continue

        # Exact match
        if norm_actual == norm_expected:
            return ct  # perfect match — return immediately

        # Containment (Claude tagged more text)
        if norm_expected in norm_actual:
            score = len(norm_expected) / len(norm_actual) + 0.5  # boost
            if score > best_score:
                best_score = score
                best_tag = ct
            continue

        # Containment (Claude tagged less text — partial)
        if norm_actual in norm_expected:
            score = len(norm_actual) / len(norm_expected) + 0.3
            if score > best_score:
                best_score = score
                best_tag = ct
            continue

        # Fuzzy match
        ratio = SequenceMatcher(None, norm_expected, norm_actual).ratio()
        if ratio >= FUZZY_THRESHOLD and ratio > best_score:
            best_score = ratio
            best_tag = ct

    return best_tag


# ---------------------------------------------------------------------------
# Core comparison
# ---------------------------------------------------------------------------

def compare_spans(
    pymupdf_report: Dict[str, Any],
    claude_text: str,
) -> Tuple[List[ValidationRow], List[FalsePositiveRow], Dict[str, Any]]:
    """
    Compare PyMuPDF formatted spans against Claude's tagged output.

    Returns
    -------
    validated : list[ValidationRow]
        One row per PyMuPDF span with match status.
    false_positives : list[FalsePositiveRow]
        Claude tags that don't correspond to any PyMuPDF span.
    summary : dict
        Aggregate statistics (counts, precision, recall, grade).
    """
    claude_tags = parse_claude_tags(claude_text)
    pymupdf_spans = pymupdf_report.get("formatted_spans", [])

    validated: List[ValidationRow] = []
    counts = {
        "total": len(pymupdf_spans),
        "matched_full": 0,
        "matched_partial": 0,
        "missed": 0,
    }

    for sp in pymupdf_spans:
        page = sp.get("page", 0)
        bill_line = sp.get("bill_line")
        text = sp.get("text", "")
        format_types = sp.get("format_types", [])
        context = sp.get("context", "")

        # Determine expected tag type from format types
        # If span has both strikethrough and underline, strikethrough wins (deletion)
        expected_tag_type = "add"
        for ft in format_types:
            mapped = _FORMAT_TO_TAG.get(ft, "add")
            if mapped == "delete":
                expected_tag_type = "delete"
                break

        expected_tag_label = "[A>...<A]" if expected_tag_type == "add" else "[D>...<D]"
        fmt_str = ", ".join(format_types)

        # Find matching Claude tag
        match = _find_best_match(text, expected_tag_type, claude_tags)

        if match is None:
            # No match found
            row = ValidationRow(
                page=page, bill_line=bill_line, text=text,
                format_types=fmt_str, expected_tag=expected_tag_label,
                context=context, tag_found="❌ No", text_match="❌ Missing",
                missing_chars=text.strip(), extra_chars="",
                severity="Critical",
                notes=f"PyMuPDF detected {fmt_str} but Claude output has no matching {expected_tag_label} tag",
            )
            counts["missed"] += 1

        else:
            match.matched = True

            # Compare text content
            norm_expected = _normalise(text)
            norm_actual = _normalise(match.text)

            if norm_expected == norm_actual:
                # Perfect match
                # Still check raw text for minor space/punctuation differences
                missing, extra = _char_diff(text.strip(), match.text.strip())
                if missing or extra:
                    tag_found = "✅ Yes"
                    text_match = "⚠️ Minor diff"
                    severity = "Minor"
                    notes_parts = []
                    if missing:
                        notes_parts.append(f"Raw text diff — missing: '{missing}'")
                    if extra:
                        notes_parts.append(f"Raw text diff — extra: '{extra}'")
                    notes = "; ".join(notes_parts)
                    counts["matched_full"] += 1
                else:
                    tag_found = "✅ Yes"
                    text_match = "✅ Full"
                    severity = "OK"
                    notes = ""
                    missing, extra = "", ""
                    counts["matched_full"] += 1
            else:
                # Partial match — figure out what's different
                missing, extra = _char_diff(text.strip(), match.text.strip())
                tag_found = "⚠️ Partial"

                if not missing and extra:
                    text_match = "⚠️ Extra text"
                    severity = "Minor"
                    notes = f"Claude tagged extra: '{extra}'"
                elif missing and not extra:
                    text_match = "⚠️ Missing text"
                    severity = "Minor" if len(missing) <= 3 else "Critical"
                    notes = f"Missing from Claude tag: '{missing}'"
                else:
                    text_match = "⚠️ Partial"
                    severity = "Minor" if len(missing) <= 3 else "Critical"
                    notes_parts = []
                    if missing:
                        notes_parts.append(f"Missing: '{missing}'")
                    if extra:
                        notes_parts.append(f"Extra: '{extra}'")
                    notes = "; ".join(notes_parts)

                counts["matched_partial"] += 1

            row = ValidationRow(
                page=page, bill_line=bill_line, text=text,
                format_types=fmt_str, expected_tag=expected_tag_label,
                context=context, tag_found=tag_found, text_match=text_match,
                missing_chars=missing, extra_chars=extra,
                severity=severity, notes=notes,
            )

        validated.append(row)

    # --- False positives: Claude tags not matched to any PyMuPDF span ---
    false_positives: List[FalsePositiveRow] = []
    for ct in claude_tags:
        if not ct.matched:
            tag_label = "[A>...<A]" if ct.tag_type == "add" else "[D>...<D]"
            fp_text = ct.text.strip()
            if not fp_text:
                continue
            false_positives.append(FalsePositiveRow(
                tag_type=tag_label,
                tagged_text=fp_text,
                position_in_output=ct.start_pos,
                notes=f"Claude tagged as {tag_label} but no matching {ct.tag_type} formatting detected by PyMuPDF",
            ))

    # --- Summary ---
    total = counts["total"]
    matched = counts["matched_full"] + counts["matched_partial"]
    recall = (matched / total * 100) if total > 0 else 100.0
    total_claude = len(claude_tags)
    true_positives = sum(1 for ct in claude_tags if ct.matched)
    precision = (true_positives / total_claude * 100) if total_claude > 0 else 100.0

    if recall >= 98 and precision >= 98:
        grade = "A+"
    elif recall >= 95 and precision >= 95:
        grade = "A"
    elif recall >= 90 and precision >= 90:
        grade = "B+"
    elif recall >= 85 and precision >= 85:
        grade = "B"
    elif recall >= 80 and precision >= 80:
        grade = "C"
    else:
        grade = "D"

    summary = {
        "total_pymupdf_spans": total,
        "total_claude_tags": total_claude,
        "matched_full": counts["matched_full"],
        "matched_partial": counts["matched_partial"],
        "missed": counts["missed"],
        "false_positives": len(false_positives),
        "precision_pct": round(precision, 1),
        "recall_pct": round(recall, 1),
        "f1_pct": round(2 * precision * recall / (precision + recall), 1) if (precision + recall) > 0 else 0.0,
        "grade": grade,
    }

    return validated, false_positives, summary


# ---------------------------------------------------------------------------
# Excel report generation
# ---------------------------------------------------------------------------

def generate_excel_report(
    validated: List[ValidationRow],
    false_positives: List[FalsePositiveRow],
    summary: Dict[str, Any],
    output_path: str | Path,
) -> Path:
    """
    Create an Excel workbook with three sheets:
      1. Span Validation
      2. False Positives
      3. Summary
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # --- Colour definitions ---
    green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    wrap_align = Alignment(wrap_text=True, vertical="top")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ================================================================
    # Sheet 1: Span Validation
    # ================================================================
    ws1 = wb.active
    ws1.title = "Span Validation"

    headers1 = [
        "Page", "Bill Line", "Text (PyMuPDF)", "Format Type",
        "Expected Tag", "Context", "Tag Found?", "Text Match",
        "Missing Characters", "Extra Characters", "Severity", "Notes",
    ]
    for col_idx, hdr in enumerate(headers1, 1):
        cell = ws1.cell(row=1, column=col_idx, value=hdr)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, row in enumerate(validated, 2):
        values = [
            row.page,
            row.bill_line if row.bill_line is not None else "N/A",
            row.text,
            row.format_types,
            row.expected_tag,
            row.context,
            row.tag_found,
            row.text_match,
            row.missing_chars,
            row.extra_chars,
            row.severity,
            row.notes,
        ]
        for col_idx, val in enumerate(values, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = wrap_align
            cell.border = thin_border

        # Conditional row colouring based on severity
        if row.severity == "OK":
            row_fill = green_fill
        elif row.severity == "Minor":
            row_fill = yellow_fill
        else:  # Critical
            row_fill = red_fill

        for col_idx in range(1, len(headers1) + 1):
            ws1.cell(row=row_idx, column=col_idx).fill = row_fill

    # Column widths
    col_widths1 = [8, 10, 40, 18, 14, 50, 14, 16, 25, 25, 12, 50]
    for i, w in enumerate(col_widths1, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # Freeze top row
    ws1.freeze_panes = "A2"
    # Auto-filter
    ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers1))}1"

    # ================================================================
    # Sheet 2: False Positives
    # ================================================================
    ws2 = wb.create_sheet("False Positives")

    headers2 = ["Tag Type", "Tagged Text", "Position in Output", "Notes"]
    for col_idx, hdr in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col_idx, value=hdr)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    for row_idx, row in enumerate(false_positives, 2):
        values = [row.tag_type, row.tagged_text, row.position_in_output, row.notes]
        for col_idx, val in enumerate(values, 1):
            cell = ws2.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = wrap_align
            cell.border = thin_border
            cell.fill = yellow_fill

    col_widths2 = [14, 60, 20, 60]
    for i, w in enumerate(col_widths2, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    ws2.freeze_panes = "A2"

    if not false_positives:
        ws2.cell(row=2, column=1, value="No false positives detected").font = Font(italic=True)

    # ================================================================
    # Sheet 3: Summary
    # ================================================================
    ws3 = wb.create_sheet("Summary")

    summary_rows = [
        ("Metric", "Value"),
        ("Total PyMuPDF Formatted Spans", summary.get("total_pymupdf_spans", 0)),
        ("Total Claude Tags", summary.get("total_claude_tags", 0)),
        ("", ""),
        ("Fully Matched", summary.get("matched_full", 0)),
        ("Partially Matched", summary.get("matched_partial", 0)),
        ("Missed (not tagged by Claude)", summary.get("missed", 0)),
        ("False Positives (Claude only)", summary.get("false_positives", 0)),
        ("", ""),
        ("Precision (%)", summary.get("precision_pct", 0)),
        ("Recall (%)", summary.get("recall_pct", 0)),
        ("F1 Score (%)", summary.get("f1_pct", 0)),
        ("Grade", summary.get("grade", "")),
    ]

    for row_idx, (metric, value) in enumerate(summary_rows, 1):
        cell_m = ws3.cell(row=row_idx, column=1, value=metric)
        cell_v = ws3.cell(row=row_idx, column=2, value=value)
        cell_m.border = thin_border
        cell_v.border = thin_border

        if row_idx == 1:
            cell_m.fill = header_fill
            cell_m.font = header_font
            cell_v.fill = header_fill
            cell_v.font = header_font
        elif metric == "Grade":
            cell_v.font = Font(bold=True, size=14)
            grade = str(value)
            if grade.startswith("A"):
                cell_v.fill = green_fill
            elif grade.startswith("B"):
                cell_v.fill = yellow_fill
            else:
                cell_v.fill = red_fill

    ws3.column_dimensions["A"].width = 35
    ws3.column_dimensions["B"].width = 15

    # ================================================================
    # Severity breakdown (mini table on Summary sheet)
    # ================================================================
    severity_counts = {"OK": 0, "Minor": 0, "Critical": 0}
    for row in validated:
        severity_counts[row.severity] = severity_counts.get(row.severity, 0) + 1

    start_row = len(summary_rows) + 2
    ws3.cell(row=start_row, column=1, value="Severity Breakdown").font = Font(bold=True)
    ws3.cell(row=start_row, column=1).border = thin_border
    ws3.cell(row=start_row, column=2).border = thin_border

    for i, (sev, count) in enumerate(severity_counts.items()):
        r = start_row + 1 + i
        cell_s = ws3.cell(row=r, column=1, value=sev)
        cell_c = ws3.cell(row=r, column=2, value=count)
        cell_s.border = thin_border
        cell_c.border = thin_border
        if sev == "OK":
            cell_s.fill = green_fill
        elif sev == "Minor":
            cell_s.fill = yellow_fill
        else:
            cell_s.fill = red_fill

    # Save
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    return output_path


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def validate_conversion(
    claude_text: str,
    pymupdf_report: Dict[str, Any],
    output_path: str | Path,
) -> Tuple[Dict[str, Any], Path]:
    """
    Run full Tier-1 validation and generate Excel report.

    Parameters
    ----------
    claude_text : str
        Claude's converted text with ``[A>…<A]`` / ``[D>…<D]`` tags.
    pymupdf_report : dict
        JSON report from ``extract_pdf_json_report()``.
    output_path : str or Path
        Path for the output ``.xlsx`` file.

    Returns
    -------
    summary : dict
        Aggregate statistics.
    excel_path : Path
        Path to the generated Excel report.
    """
    validated, false_positives, summary = compare_spans(pymupdf_report, claude_text)
    excel_path = generate_excel_report(validated, false_positives, summary, output_path)
    return summary, excel_path


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Validate Claude tagged output against PyMuPDF JSON report"
    )
    parser.add_argument("claude_txt", help="Path to Claude's converted .txt file")
    parser.add_argument("pymupdf_json", help="Path to PyMuPDF JSON report")
    parser.add_argument(
        "-o", "--output", default=None,
        help="Output Excel path (default: <claude_txt>_validation.xlsx)"
    )
    args = parser.parse_args()

    claude_path = Path(args.claude_txt)
    report_path = Path(args.pymupdf_json)

    if not claude_path.exists():
        print(f"❌ Claude output not found: {claude_path}")
        sys.exit(1)
    if not report_path.exists():
        print(f"❌ PyMuPDF report not found: {report_path}")
        sys.exit(1)

    claude_text = claude_path.read_text(encoding="utf-8")
    pymupdf_report = json.loads(report_path.read_text(encoding="utf-8"))

    output = args.output or str(claude_path).replace(".txt", "_validation.xlsx")

    print(f"📄 Claude output:  {claude_path}")
    print(f"📄 PyMuPDF report: {report_path}")
    print(f"📄 Output:         {output}")
    print()

    summary, excel_path = validate_conversion(claude_text, pymupdf_report, output)

    print(f"✅ Validation report: {excel_path}")
    print()
    print(f"   PyMuPDF spans:     {summary['total_pymupdf_spans']}")
    print(f"   Claude tags:       {summary['total_claude_tags']}")
    print(f"   Fully matched:     {summary['matched_full']}")
    print(f"   Partially matched: {summary['matched_partial']}")
    print(f"   Missed:            {summary['missed']}")
    print(f"   False positives:   {summary['false_positives']}")
    print(f"   Precision:         {summary['precision_pct']}%")
    print(f"   Recall:            {summary['recall_pct']}%")
    print(f"   F1 Score:          {summary['f1_pct']}%")
    print(f"   Grade:             {summary['grade']}")


if __name__ == "__main__":
    main()
