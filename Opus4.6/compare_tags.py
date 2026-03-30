"""
Char-by-char comparison of ALL tagged content:
  PyMuPDF (LLM-cleaned text) vs Claude (converted output)

Pure text-to-text — no JSON report, no bill_line filtering.
Extracts [A>...<A] tags and [brackets] / [D>...<D] directly from both files.
Works universally for any PDF format.

Usage:
  As module:   from compare_tags import generate_tag_comparison_excel
  As script:   python compare_tags.py <pymupdf_clean.txt> <claude_converted.txt> <output.xlsx>
"""
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Regex patterns ──────────────────────────────────────────────────
_A_RE = re.compile(r"\[A>(.*?)<A\]", re.DOTALL)
_D_RE = re.compile(r"\[D>(.*?)<D\]", re.DOTALL)
_BRACKET_RE = re.compile(
    r"(?<!\[A>)(?<!\[D>)\[(?!A>|D>)([^\[\]]+?)\](?!<[AD])", re.DOTALL
)

# ── Styles ──────────────────────────────────────────────────────────
_HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor="2F5496")
_OK_FILL     = PatternFill("solid", fgColor="C6EFCE")
_DIFF_FILL   = PatternFill("solid", fgColor="FFC7CE")
_MISS_FILL   = PatternFill("solid", fgColor="FFD966")
_EXTRA_FILL  = PatternFill("solid", fgColor="D9E2F3")
_WRAP        = Alignment(wrap_text=True, vertical="top")
_THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


# ── Helpers ─────────────────────────────────────────────────────────
def _norm(t: str) -> str:
    return " ".join(t.split())


def _char_diff_positions(s1: str, s2: str):
    """Return list of (pos, char_in_s1, char_in_s2) tuples."""
    if s1 == s2:
        return []
    result = []
    for i in range(max(len(s1), len(s2))):
        c1 = s1[i] if i < len(s1) else ""
        c2 = s2[i] if i < len(s2) else ""
        if c1 != c2:
            result.append((i, c1 or "<END>", c2 or "<END>"))
    return result


def _match_tags(py_list, cl_list):
    """Match PyMuPDF tags to Claude tags by content."""
    used = set()
    results = []
    for i, py_text in enumerate(py_list):
        best_j = None
        best_score = 0
        for j, cl_text in enumerate(cl_list):
            if j in used:
                continue
            if py_text == cl_text:
                best_j = j
                best_score = 1.0
                break
            if py_text in cl_text or cl_text in py_text:
                denom = max(len(py_text), len(cl_text))
                overlap = min(len(py_text), len(cl_text)) / denom if denom > 0 else 0
                if overlap > best_score:
                    best_j = j
                    best_score = overlap
        if best_j is not None:
            used.add(best_j)
            cl_text = cl_list[best_j]
            diffs = _char_diff_positions(py_text, cl_text)
            results.append({
                "py_idx": i, "cl_idx": best_j,
                "py_text": py_text, "cl_text": cl_text,
                "match": len(diffs) == 0,
                "diff_count": len(diffs),
                "diffs": diffs,
                "py_len": len(py_text), "cl_len": len(cl_text),
            })
        else:
            results.append({
                "py_idx": i, "cl_idx": None,
                "py_text": py_text, "cl_text": "",
                "match": False, "diff_count": -1,
                "diffs": [], "py_len": len(py_text), "cl_len": 0,
            })
    extras = []
    for j in range(len(cl_list)):
        if j not in used:
            extras.append({"cl_idx": j, "cl_text": cl_list[j]})
    return results, extras


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _THIN_BORDER


def _style_row(ws, row, ncols, fill=None):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.alignment = _WRAP
        cell.border = _THIN_BORDER
        if fill:
            cell.fill = fill


# ════════════════════════════════════════════════════════════════════
# Public API
# ════════════════════════════════════════════════════════════════════

def generate_tag_comparison_excel(
    pymupdf_text: str,
    claude_text: str,
    excel_path: str,
) -> dict:
    """
    Compare ALL [A>…<A] and [D>…<D] / [bracket] tags between
    PyMuPDF-cleaned text and Claude-converted text, write a styled
    Excel report, and return a summary dict.

    Parameters
    ----------
    pymupdf_text : str
        LLM-cleaned + consolidated PyMuPDF output.
    claude_text : str
        Claude-converted text (with [A>…<A] and [D>…<D] tags).
    excel_path : str
        Destination path for the Excel report.

    Returns
    -------
    dict with keys:
        pymupdf_additions, claude_additions, pymupdf_deletions,
        claude_deletions, total_compared, exact_matches, differences,
        match_rate_pct, diff_groups, excel_path
    """
    # ── Extract tags ────────────────────────────────────────────────
    py_additions = [_norm(t) for t in _A_RE.findall(pymupdf_text)]
    py_deletions = [_norm(t) for t in _BRACKET_RE.findall(pymupdf_text)]
    cl_a_norm    = [_norm(t) for t in _A_RE.findall(claude_text)]

    cl_d_clean = []
    for d in _D_RE.findall(claude_text):
        d = d.strip()
        if d.startswith("[") and d.endswith("]"):
            d = d[1:-1]
        cl_d_clean.append(_norm(d))

    # ── Match ───────────────────────────────────────────────────────
    add_results, add_extras = _match_tags(py_additions, cl_a_norm)
    del_results, del_extras = _match_tags(py_deletions, cl_d_clean)

    total_ok   = sum(1 for r in add_results + del_results if r["match"])
    total_diff = sum(1 for r in add_results + del_results if not r["match"])
    total      = len(add_results) + len(del_results)

    # ── Build Excel ─────────────────────────────────────────────────
    wb = Workbook()

    # ──── Sheet 1: Tag Comparison Detail ────────────────────────────
    ws1 = wb.active
    ws1.title = "Tag Comparison"
    headers = [
        "#", "Type",
        "PyMuPDF Text (Ground Truth)", "Claude Text (Converted)",
        "Status", "Char Diffs", "PyMuPDF Len", "Claude Len", "Len Diff",
        "First Diff Detail",
    ]
    NCOLS = len(headers)

    # Title
    ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NCOLS)
    title_cell = ws1.cell(
        row=1, column=1,
        value="Char-by-Char Tag Comparison: PyMuPDF (Ground Truth) vs Claude (Converted)",
    )
    title_cell.font = Font(name="Calibri", bold=True, size=14, color="2F5496")
    title_cell.alignment = Alignment(horizontal="center")

    # Subtitle
    ws1.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOLS)
    ws1.cell(
        row=2, column=1,
        value=(
            f"Additions: {len(py_additions)} PyMuPDF / {len(cl_a_norm)} Claude  |  "
            f"Deletions: {len(py_deletions)} PyMuPDF / {len(cl_d_clean)} Claude  |  "
            f"Result: {total_ok}/{total} exact match, {total_diff} with differences"
        ),
    ).font = Font(name="Calibri", size=11, italic=True)

    # Headers
    for c, h in enumerate(headers, 1):
        ws1.cell(row=4, column=c, value=h)
    _style_header(ws1, 4, NCOLS)

    # Helper to write one result row
    def _write_result(ws, r, num, tag_type, res):
        status = "Exact Match" if res["match"] else ("No Match" if res["cl_idx"] is None else "Difference")
        fill = _OK_FILL if res["match"] else (_MISS_FILL if res["cl_idx"] is None else _DIFF_FILL)
        detail = ""
        if res["diffs"]:
            parts = [f"pos {p}: '{c1}' -> '{c2}'" for p, c1, c2 in res["diffs"][:5]]
            detail = "; ".join(parts)
            if len(res["diffs"]) > 5:
                detail += f" ... +{len(res['diffs'])-5} more"
        ws.cell(row=r, column=1, value=num)
        ws.cell(row=r, column=2, value=tag_type)
        ws.cell(row=r, column=3, value=res["py_text"])
        ws.cell(row=r, column=4, value=res["cl_text"])
        ws.cell(row=r, column=5, value=status)
        ws.cell(row=r, column=6, value=res["diff_count"] if res["diff_count"] >= 0 else "N/A")
        ws.cell(row=r, column=7, value=res["py_len"])
        ws.cell(row=r, column=8, value=res["cl_len"])
        ws.cell(row=r, column=9, value=res["cl_len"] - res["py_len"] if res["cl_idx"] is not None else "N/A")
        ws.cell(row=r, column=10, value=detail)
        _style_row(ws, r, NCOLS, fill)

    def _write_extra(ws, r, num, tag_type, cl_text):
        ws.cell(row=r, column=1, value=num)
        ws.cell(row=r, column=2, value=tag_type)
        ws.cell(row=r, column=3, value="")
        ws.cell(row=r, column=4, value=cl_text)
        ws.cell(row=r, column=5, value="Extra in Claude")
        _style_row(ws, r, NCOLS, _EXTRA_FILL)

    r, num = 5, 0
    for res in add_results:
        num += 1; _write_result(ws1, r, num, "Addition [A>]", res); r += 1
    for ex in add_extras:
        num += 1; _write_extra(ws1, r, num, "Addition [A>]", ex["cl_text"]); r += 1
    for res in del_results:
        num += 1; _write_result(ws1, r, num, "Deletion [D>]", res); r += 1
    for ex in del_extras:
        num += 1; _write_extra(ws1, r, num, "Deletion [D>]", ex["cl_text"]); r += 1

    col_widths = [5, 14, 55, 55, 14, 10, 10, 10, 9, 45]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w
    ws1.freeze_panes = "A5"
    if r > 5:
        ws1.auto_filter.ref = f"A4:{get_column_letter(NCOLS)}{r-1}"

    # ──── Sheet 2: Char Diff Detail ─────────────────────────────────
    ws2 = wb.create_sheet("Char Diffs Detail")
    h2 = ["#", "Type", "Position", "PyMuPDF Char", "Claude Char",
           "PyMuPDF Full Text", "Claude Full Text"]
    NC2 = len(h2)
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NC2)
    ws2.cell(row=1, column=1,
             value="Character-Level Differences (only mismatched tags)").font = Font(
        name="Calibri", bold=True, size=14, color="2F5496")
    row2 = 3
    for c, h in enumerate(h2, 1):
        ws2.cell(row=row2, column=c, value=h)
    _style_header(ws2, row2, NC2)
    row2 += 1

    diff_num = 0
    for res, tag_type in (
        [(r, "Addition") for r in add_results] +
        [(r, "Deletion") for r in del_results]
    ):
        if res["match"] or not res["diffs"]:
            continue
        diff_num += 1
        first_row = row2
        for pos, c1, c2 in res["diffs"]:
            ws2.cell(row=row2, column=1, value=diff_num)
            ws2.cell(row=row2, column=2, value=tag_type)
            ws2.cell(row=row2, column=3, value=pos)
            ws2.cell(row=row2, column=4, value=c1)
            ws2.cell(row=row2, column=5, value=c2)
            if row2 == first_row:
                ws2.cell(row=row2, column=6, value=res["py_text"])
                ws2.cell(row=row2, column=7, value=res["cl_text"])
            _style_row(ws2, row2, NC2, _DIFF_FILL)
            row2 += 1

    for i, w in enumerate([5, 10, 9, 14, 14, 55, 55], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A4"

    # ──── Sheet 3: Summary ──────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")
    ws3.merge_cells("A1:D1")
    ws3.cell(row=1, column=1, value="Tag Comparison Summary").font = Font(
        name="Calibri", bold=True, size=14, color="2F5496")

    match_rate = f"{total_ok/total*100:.1f}%" if total > 0 else "N/A"
    summary_rows = [
        ("", ""),
        ("ADDITIONS [A>...<A]", ""),
        ("PyMuPDF tags", len(py_additions)),
        ("Claude tags", len(cl_a_norm)),
        ("Exact matches", sum(1 for r in add_results if r["match"])),
        ("With char differences", sum(1 for r in add_results if not r["match"] and r["cl_idx"] is not None)),
        ("Unmatched (PyMuPDF only)", sum(1 for r in add_results if r["cl_idx"] is None)),
        ("Extra (Claude only)", len(add_extras)),
        ("", ""),
        ("DELETIONS [D>...<D]", ""),
        ("PyMuPDF bracket tags", len(py_deletions)),
        ("Claude tags", len(cl_d_clean)),
        ("Exact matches", sum(1 for r in del_results if r["match"])),
        ("With char differences", sum(1 for r in del_results if not r["match"] and r["cl_idx"] is not None)),
        ("Unmatched (PyMuPDF only)", sum(1 for r in del_results if r["cl_idx"] is None)),
        ("Extra (Claude only)", len(del_extras)),
        ("", ""),
        ("GRAND TOTAL", ""),
        ("Total tags compared", total),
        ("Exact char matches", total_ok),
        ("With differences", total_diff),
        ("Match rate", match_rate),
    ]
    for i, (label, val) in enumerate(summary_rows, 3):
        c1 = ws3.cell(row=i, column=1, value=label)
        c2 = ws3.cell(row=i, column=2, value=val)
        if label and not val and label != "":
            c1.font = Font(bold=True, size=12)
        c1.alignment = _WRAP
        c2.alignment = _WRAP
    ws3.column_dimensions["A"].width = 30
    ws3.column_dimensions["B"].width = 15

    # ── Save ────────────────────────────────────────────────────────
    Path(excel_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(excel_path)

    return {
        "pymupdf_additions": len(py_additions),
        "claude_additions": len(cl_a_norm),
        "pymupdf_deletions": len(py_deletions),
        "claude_deletions": len(cl_d_clean),
        "total_compared": total,
        "exact_matches": total_ok,
        "differences": total_diff,
        "match_rate_pct": round(total_ok / total * 100, 1) if total > 0 else 0,
        "diff_groups": diff_num,
        "excel_path": str(excel_path),
    }


# ════════════════════════════════════════════════════════════════════
# Highlighted HTML — inline diff on Claude text
# ════════════════════════════════════════════════════════════════════

def generate_highlighted_html(
    pymupdf_text: str,
    claude_text: str,
    html_path: str,
) -> dict:
    """
    Produce an HTML rendering of the **Claude converted text** where every
    ``[A>…<A]`` and ``[D>…<D]`` tag is colour-coded against the PyMuPDF
    ground truth:

    * **Green** — exact char match with PyMuPDF  (auto-approved)
    * **Red**   — difference found (hover shows PyMuPDF version + char diffs)
    * **Yellow** — tag only in Claude (no PyMuPDF match)
    * **Grey**  — tag only in PyMuPDF (missing from Claude) listed at bottom

    Returns the same summary dict as ``generate_tag_comparison_excel``.
    """
    import html as _html

    # ── Extract & match (reuse same logic) ──────────────────────────
    py_additions = [_norm(t) for t in _A_RE.findall(pymupdf_text)]
    py_deletions = [_norm(t) for t in _BRACKET_RE.findall(pymupdf_text)]
    cl_a_norm    = [_norm(t) for t in _A_RE.findall(claude_text)]

    cl_d_clean = []
    for d in _D_RE.findall(claude_text):
        d = d.strip()
        if d.startswith("[") and d.endswith("]"):
            d = d[1:-1]
        cl_d_clean.append(_norm(d))

    add_results, add_extras = _match_tags(py_additions, cl_a_norm)
    del_results, del_extras = _match_tags(py_deletions, cl_d_clean)

    # Build lookup: claude_tag_index → result for additions
    cl_add_lookup = {}   # cl_idx → result
    for res in add_results:
        if res["cl_idx"] is not None:
            cl_add_lookup[res["cl_idx"]] = res
    # additions with no Claude match
    py_only_adds = [res for res in add_results if res["cl_idx"] is None]

    cl_del_lookup = {}
    for res in del_results:
        if res["cl_idx"] is not None:
            cl_del_lookup[res["cl_idx"]] = res
    py_only_dels = [res for res in del_results if res["cl_idx"] is None]

    # ── Walk Claude text and replace tags with highlighted spans ────
    # We process [A>…<A] and [D>…<D] tags in document order
    _TAG_RE = re.compile(r"(\[A>.*?<A\]|\[D>.*?<D\])", re.DOTALL)

    add_counter = 0  # sequential index into cl_a_norm
    del_counter = 0  # sequential index into cl_d_clean

    parts = []
    last_end = 0
    for m in _TAG_RE.finditer(claude_text):
        # Plain text before this tag
        plain = claude_text[last_end:m.start()]
        parts.append(_html.escape(plain))

        raw_tag = m.group(0)
        is_addition = raw_tag.startswith("[A>")

        if is_addition:
            idx = add_counter
            add_counter += 1
            res = cl_add_lookup.get(idx)
            if res is None:
                # extra in Claude, no PyMuPDF match
                css = "background:#FFD966;border-bottom:2px solid #E6A800;"
                title = "EXTRA — not in PyMuPDF ground truth"
                icon = "&#x26A0;"  # ⚠
            elif res["match"]:
                css = "background:#C6EFCE;border-bottom:2px solid #27AE60;"
                title = "EXACT MATCH with PyMuPDF"
                icon = "&#x2714;"  # ✔
            else:
                css = "background:#FFC7CE;border-bottom:2px solid #E74C3C;"
                diff_info = _describe_diff(res)
                diff_detail = diff_info["tooltip"]
                action = diff_info["action"]
                title = f"DIFFERENCE \u2014 PyMuPDF: {_html.escape(res['py_text'])}\\n{diff_detail}"
                icon = "&#x2716;"  # \u2716
            # Render the tag content (without [A>\u2026<A] markers)
            inner = _A_RE.findall(raw_tag)[0] if _A_RE.findall(raw_tag) else raw_tag
            action_badge = ""
            if not (res and res.get("match", True)) and res is not None and res.get("cl_idx") is not None:
                action_badge = f'<span class="action-badge">{_html.escape(action)}</span>'
            parts.append(
                f'<span class="tag add" style="{css}" title="{title}">'
                f'<sup class="icon">{icon}</sup>'
                f'[A&gt;{_html.escape(inner)}&lt;A]'
                f'{action_badge}'
                f'</span>'
            )
        else:
            # Deletion [D>…<D]
            idx = del_counter
            del_counter += 1
            res = cl_del_lookup.get(idx)
            if res is None:
                css = "background:#FFD966;border-bottom:2px dashed #E6A800;"
                title = "EXTRA — not in PyMuPDF ground truth"
                icon = "&#x26A0;"
            elif res["match"]:
                css = "background:#D5F5E3;border-bottom:2px dashed #27AE60;"
                title = "EXACT MATCH with PyMuPDF"
                icon = "&#x2714;"
            else:
                css = "background:#FADBD8;border-bottom:2px dashed #E74C3C;"
                diff_info = _describe_diff(res)
                diff_detail = diff_info["tooltip"]
                action = diff_info["action"]
                title = f"DIFFERENCE \u2014 PyMuPDF: {_html.escape(res['py_text'])}\\n{diff_detail}"
                icon = "&#x2716;"
            inner = _D_RE.findall(raw_tag)[0] if _D_RE.findall(raw_tag) else raw_tag
            action_badge = ""
            if not (res and res.get("match", True)) and res is not None and res.get("cl_idx") is not None:
                action_badge = f'<span class="action-badge">{_html.escape(action)}</span>'
            parts.append(
                f'<span class="tag del" style="{css}" title="{title}">'
                f'<sup class="icon">{icon}</sup>'
                f'[D&gt;{_html.escape(inner)}&lt;D]'
                f'{action_badge}'
                f'</span>'
            )
        last_end = m.end()

    # Remaining plain text
    parts.append(_html.escape(claude_text[last_end:]))

    body_html = "".join(parts)

    # ── Build missing tags section (in PyMuPDF but not Claude) ──────
    missing_html = ""
    if py_only_adds or py_only_dels:
        missing_items = []
        for res in py_only_adds:
            missing_items.append(f'<li class="miss-add">[A&gt;] {_html.escape(res["py_text"])}</li>')
        for res in py_only_dels:
            missing_items.append(f'<li class="miss-del">[D&gt;] {_html.escape(res["py_text"])}</li>')
        missing_html = (
            '<div class="missing-section">'
            '<h3>Tags in PyMuPDF but NOT in Claude (potential misses)</h3>'
            '<ul>' + "\n".join(missing_items) + '</ul>'
            '</div>'
        )

    # ── Stats ───────────────────────────────────────────────────────
    total_ok   = sum(1 for r in add_results + del_results if r["match"])
    total_diff = sum(1 for r in add_results + del_results if not r["match"])
    total      = len(add_results) + len(del_results)
    match_pct  = round(total_ok / total * 100, 1) if total > 0 else 0

    # ── Assemble full HTML ──────────────────────────────────────────
    full_html = _HTML_TEMPLATE.format(
        title="Claude Converted Text — Validation Highlights",
        match_pct=match_pct,
        exact=total_ok,
        diffs=total_diff,
        total=total,
        n_add_py=len(py_additions),
        n_add_cl=len(cl_a_norm),
        n_del_py=len(py_deletions),
        n_del_cl=len(cl_d_clean),
        n_extra_add=len(add_extras),
        n_extra_del=len(del_extras),
        n_miss_add=len(py_only_adds),
        n_miss_del=len(py_only_dels),
        body=body_html,
        missing=missing_html,
    )

    Path(html_path).parent.mkdir(parents=True, exist_ok=True)
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(full_html)

    return {
        "pymupdf_additions": len(py_additions),
        "claude_additions": len(cl_a_norm),
        "pymupdf_deletions": len(py_deletions),
        "claude_deletions": len(cl_d_clean),
        "total_compared": total,
        "exact_matches": total_ok,
        "differences": total_diff,
        "match_rate_pct": match_pct,
        "html_path": str(html_path),
    }


def _build_diff_title(res):
    """Build a short diff description for tooltip."""
    return _describe_diff(res)["tooltip"]


def _describe_diff(res):
    """
    Analyse the char-level diffs and return a dict with:
      action  — short actionable label shown inline  (e.g. "DELETE ','")
      tooltip — full detail for hover
      severity — 'critical' | 'minor' | 'cosmetic'
    """
    diffs = res["diffs"]
    py = res["py_text"]
    cl = res["cl_text"]

    if not diffs:
        return {"action": "", "tooltip": "", "severity": "cosmetic"}

    # Classify: extra chars at end of Claude
    extra_at_end = all(c1 == "<END>" for _, c1, c2 in diffs)
    # Classify: missing chars at end of Claude
    missing_at_end = all(c2 == "<END>" for _, c1, c2 in diffs)
    # Classify: extra chars at start of Claude (leading insertion)
    extra_at_start = len(cl) > len(py) and cl.endswith(py[-min(len(py),20):]) if py else False

    # ── Extra at end ────────────────────────────────────────────────
    if extra_at_end:
        extra_chars = "".join(c2 for _, _, c2 in diffs)
        extra_repr = repr(extra_chars)
        if extra_chars.strip() == "":
            return {
                "action": f"DELETE trailing whitespace",
                "tooltip": f"Claude has extra trailing whitespace ({len(extra_chars)} char(s))",
                "severity": "cosmetic",
            }
        if all(c in ",.;:" for c in extra_chars):
            return {
                "action": f"DELETE trailing {extra_repr}",
                "tooltip": f"Claude has extra trailing punctuation {extra_repr} — not in PyMuPDF ground truth",
                "severity": "minor",
            }
        return {
            "action": f"DELETE extra {extra_repr} at end",
            "tooltip": f"Claude has {len(extra_chars)} extra char(s) at end: {extra_repr}\\nPyMuPDF ends: ...{py[-20:]}",
            "severity": "minor",
        }

    # ── Missing at end ──────────────────────────────────────────────
    if missing_at_end:
        missing_chars = "".join(c1 for _, c1, _ in diffs)
        missing_repr = repr(missing_chars)
        if missing_chars.strip() == "":
            return {
                "action": f"ADD trailing whitespace",
                "tooltip": f"Claude is missing trailing whitespace ({len(missing_chars)} char(s))",
                "severity": "cosmetic",
            }
        return {
            "action": f"ADD missing {missing_repr} at end",
            "tooltip": f"Claude is missing {len(missing_chars)} char(s) at end: {missing_repr}\\nPyMuPDF: ...{py[-30:]}",
            "severity": "critical" if len(missing_chars) > 3 else "minor",
        }

    # ── Extra at start ──────────────────────────────────────────────
    if extra_at_start:
        leading = cl[:len(cl) - len(py)]
        leading_repr = repr(leading)
        if leading.strip() in ",.;:":
            return {
                "action": f"DELETE leading {leading_repr}",
                "tooltip": f"Claude has extra leading punctuation {leading_repr} — not in PyMuPDF\\nPyMuPDF: {py[:40]}",
                "severity": "minor",
            }
        return {
            "action": f"DELETE extra {leading_repr} at start",
            "tooltip": f"Claude has {len(leading)} extra leading char(s): {leading_repr}\\nPyMuPDF starts: {py[:40]}",
            "severity": "critical" if len(leading) > 3 else "minor",
        }

    # ── General: word-level diff ────────────────────────────────────
    py_words = py.split()
    cl_words = cl.split()
    if py_words != cl_words:
        # Find first differing word
        for i, (pw, cw) in enumerate(zip(py_words, cl_words)):
            if pw != cw:
                return {
                    "action": f"CHANGE '{cw}' to '{pw}'",
                    "tooltip": f"Word {i+1} differs — Claude: '{cw}' vs PyMuPDF: '{pw}'\\nFull PyMuPDF: {py[:80]}",
                    "severity": "critical",
                }
        # Different lengths
        if len(cl_words) > len(py_words):
            extra = " ".join(cl_words[len(py_words):])
            return {
                "action": f"DELETE extra words: '{extra[:30]}'",
                "tooltip": f"Claude has {len(cl_words)-len(py_words)} extra word(s): '{extra}'",
                "severity": "critical",
            }
        else:
            missing = " ".join(py_words[len(cl_words):])
            return {
                "action": f"ADD missing words: '{missing[:30]}'",
                "tooltip": f"Claude is missing {len(py_words)-len(cl_words)} word(s): '{missing}'",
                "severity": "critical",
            }

    # ── Fallback: char-level description ────────────────────────────
    n = len(diffs)
    parts = []
    for pos, c1, c2 in diffs[:3]:
        if c1 == "<END>":
            parts.append(f"DELETE extra '{c2}' at pos {pos}")
        elif c2 == "<END>":
            parts.append(f"ADD missing '{c1}' at pos {pos}")
        else:
            parts.append(f"CHANGE '{c2}' to '{c1}' at pos {pos}")
    action = parts[0] if len(parts) == 1 else f"{n} char diffs"
    tooltip = "; ".join(parts)
    if n > 3:
        tooltip += f" +{n-3} more"
    return {
        "action": action,
        "tooltip": f"{tooltip}\\nPyMuPDF: {py[:80]}",
        "severity": "minor" if n <= 2 else "critical",
    }


_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{
    font-family: 'Segoe UI', Consolas, monospace;
    font-size: 14px;
    line-height: 1.7;
    max-width: 1100px;
    margin: 30px auto;
    padding: 0 20px;
    background: #FAFAFA;
    color: #222;
  }}
  h1 {{ color: #2F5496; font-size: 20px; }}
  .stats {{
    background: #EBF5FB;
    border-left: 4px solid #2F5496;
    padding: 12px 18px;
    margin-bottom: 20px;
    font-size: 13px;
  }}
  .stats b {{ color: #2F5496; }}
  .legend {{
    display: flex; gap: 20px; flex-wrap: wrap;
    margin-bottom: 18px; font-size: 13px;
  }}
  .legend span {{
    padding: 3px 10px; border-radius: 3px; font-weight: 600;
  }}
  .leg-ok   {{ background: #C6EFCE; border: 1px solid #27AE60; }}
  .leg-diff {{ background: #FFC7CE; border: 1px solid #E74C3C; }}
  .leg-xtra {{ background: #FFD966; border: 1px solid #E6A800; }}
  .leg-miss {{ background: #D5D8DC; border: 1px solid #888; }}

  .bill-text {{
    white-space: pre-wrap;
    word-wrap: break-word;
    background: #fff;
    border: 1px solid #ddd;
    padding: 20px;
    border-radius: 4px;
    box-shadow: 0 1px 3px rgba(0,0,0,0.08);
  }}
  .tag {{
    padding: 1px 2px;
    border-radius: 3px;
    cursor: help;
    position: relative;
  }}
  .icon {{
    font-size: 10px;
    margin-right: 1px;
  }}
  .action-badge {{
    display: inline-block;
    background: #C0392B;
    color: #fff;
    font-size: 11px;
    font-weight: 700;
    font-family: 'Segoe UI', sans-serif;
    padding: 2px 7px;
    border-radius: 3px;
    margin-left: 4px;
    vertical-align: middle;
    white-space: nowrap;
  }}
  .missing-section {{
    margin-top: 30px;
    background: #F9EBEA;
    border: 1px solid #E74C3C;
    border-radius: 4px;
    padding: 15px;
  }}
  .missing-section h3 {{ color: #C0392B; margin-top: 0; }}
  .miss-add {{ color: #1A5276; }}
  .miss-del {{ color: #7B241C; }}
  /* tooltip on hover */
  .tag:hover::after {{
    content: attr(title);
    position: absolute;
    bottom: 110%;
    left: 0;
    background: #333;
    color: #fff;
    padding: 6px 10px;
    border-radius: 4px;
    font-size: 12px;
    white-space: pre;
    max-width: 600px;
    z-index: 999;
    box-shadow: 0 2px 8px rgba(0,0,0,0.3);
  }}
  @media print {{
    .tag {{ border: 1px solid #999 !important; }}
  }}
</style>
</head>
<body>
<h1>{title}</h1>

<div class="stats">
  <b>Match Rate: {match_pct}%</b> &mdash;
  {exact}/{total} exact match, {diffs} with differences<br>
  Additions: {n_add_py} PyMuPDF / {n_add_cl} Claude &nbsp;|&nbsp;
  Deletions: {n_del_py} PyMuPDF / {n_del_cl} Claude<br>
  Extra in Claude: {n_extra_add} add + {n_extra_del} del &nbsp;|&nbsp;
  Missing from Claude: {n_miss_add} add + {n_miss_del} del
</div>

<div class="legend">
  <span class="leg-ok">&#x2714; Exact Match</span>
  <span class="leg-diff">&#x2716; Difference (hover for detail)</span>
  <span class="leg-xtra">&#x26A0; Extra in Claude</span>
  <span class="leg-miss">&#x25CB; Missing from Claude</span>
</div>

<div class="bill-text">{body}</div>

{missing}

</body>
</html>"""


# ── Standalone entry point ──────────────────────────────────────────
if __name__ == "__main__":
    import sys

    if len(sys.argv) >= 3:
        pymupdf_file = sys.argv[1]
        claude_file  = sys.argv[2]
        excel_out    = sys.argv[3] if len(sys.argv) > 3 else "tag_comparison.xlsx"
    else:
        # Default paths (backward-compatible)
        pymupdf_file = "2025_2025_7_3_2_001193_0_4_1_20250130_0_converted_pymupdf_clean.txt"
        claude_file  = "2025_2025_7_3_2_001193_0_4_1_20250130_0_converted.txt"
        excel_out    = r"..\tmp\ct_tag_comparison.xlsx"

    with open(pymupdf_file, "r", encoding="utf-8") as f:
        pymupdf_text = f.read()
    with open(claude_file, "r", encoding="utf-8") as f:
        claude_text = f.read()

    summary = generate_tag_comparison_excel(pymupdf_text, claude_text, excel_out)

    print(f"Excel saved to: {summary['excel_path']}")
    print(f"  Sheet 1: Tag Comparison — {summary['total_compared']} tag pairs + extras")
    print(f"  Sheet 2: Char Diffs Detail — {summary['diff_groups']} diff group(s)")
    print(f"  Sheet 3: Summary — {summary['exact_matches']}/{summary['total_compared']} exact, "
          f"{summary['differences']} diffs ({summary['match_rate_pct']}%)")
